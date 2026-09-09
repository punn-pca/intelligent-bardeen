import openpyxl
import json

def parse_august_stock():
    wb = openpyxl.load_workbook('decrypted_stock.xlsm', data_only=True)
    print(f"Total Sheets in 'สต๊อค 2026 เดือน 8 สิงหาคม.xlsm': {len(wb.sheetnames)}")
    print(f"Sheet Names: {wb.sheetnames}\n")

    summary_data = {}

    for sname in wb.sheetnames:
        sheet = wb[sname]
        rows = list(sheet.iter_rows(values_only=True))
        non_empty = [r for r in rows if any(r)]

        print(f"=== Sheet: {sname} (Total Rows: {len(rows)}, Non-empty: {len(non_empty)}) ===")
        # Print top 10 rows
        for idx, row in enumerate(non_empty[:10]):
            clean_row = [str(c).strip() if c is not None else '' for c in row if c is not None]
            if any(clean_row):
                print(f"Row {idx+1}: {clean_row[:8]}")

        summary_data[sname] = {
            'totalRows': len(rows),
            'nonEmptyRows': len(non_empty),
            'sampleRows': [[str(c) if c is not None else '' for c in r[:10]] for r in non_empty[:5]]
        }
        print("-" * 60)

    with open('august_stock_summary.json', 'w', encoding='utf-8') as f:
        json.dump(summary_data, f, ensure_ascii=False, indent=2)

if __name__ == '__main__':
    parse_august_stock()
